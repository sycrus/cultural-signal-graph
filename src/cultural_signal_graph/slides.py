"""Slide-deck (.pptx) export of the brief (issue #16)."""

from __future__ import annotations

from pptx import Presentation

from .models import Brief


def _add_bullets(text_frame, items: list[str]) -> None:
    first = True
    for text in items:
        paragraph = text_frame.paragraphs[0] if first else text_frame.add_paragraph()
        paragraph.text = text
        first = False


def build_presentation(brief: Brief) -> Presentation:
    prs = Presentation()

    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = f"Trend Forecast — {brief.category}"
    title.placeholders[1].text = f"as of {brief.as_of}"

    if brief.claims:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Trends"
        items: list[str] = []
        for claim in brief.claims:
            sources = "; ".join(source.uri for source in claim.sources)
            items.append(f"{claim.statement} — confidence {claim.confidence}")
            items.append(f"    sources: {sources}")
        _add_bullets(slide.placeholders[1].text_frame, items)

    if brief.flags:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Conflicts & Flags"
        _add_bullets(
            slide.placeholders[1].text_frame,
            [f"{f.statement} — {f.reason} (confidence {f.confidence})" for f in brief.flags],
        )

    if brief.gated_edges:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Gated Cross-Domain Edges"
        _add_bullets(
            slide.placeholders[1].text_frame,
            [
                f"{e.left_entity} ≈ {e.right_entity} — {e.status} (confidence {e.confidence})"
                for e in brief.gated_edges
            ],
        )

    return prs


def write_presentation(brief: Brief, path) -> None:
    build_presentation(brief).save(str(path))
