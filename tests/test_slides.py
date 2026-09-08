"""Issue #16 acceptance tests for the .pptx export."""

from pptx import Presentation

from cultural_signal_graph.fixtures import ALIGNED_EDGES, DRAFTED_REASONING, EXTRACTED, SEED_CORPUS
from cultural_signal_graph.pipeline import run
from cultural_signal_graph.slides import write_presentation


def _forecast():
    return run(
        SEED_CORPUS,
        category="tech",
        extract_fn=lambda _: EXTRACTED,
        align_fn=lambda _: ALIGNED_EDGES,
        draft_fn=lambda extraction, category: DRAFTED_REASONING,
    )


def test_brief_renders_to_pptx(tmp_path):
    path = tmp_path / "brief.pptx"
    write_presentation(_forecast().brief, path)
    assert path.exists() and path.stat().st_size > 0

    prs = Presentation(str(path))
    text = "\n".join(
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "Trend Forecast" in text
    assert "Quark" in text  # a claim
    assert "solutions looking for a problem" in text  # the conflict flag
    assert "proposed" in text  # the gated cross-domain edge
